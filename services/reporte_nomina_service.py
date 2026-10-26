from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from services.nomina_service import report_data
from services.recibo_nomina_service import _fecha_larga, _money

COLUMNAS = [
    "Nómina",
    "Empleado",
    "Documento",
    "Periodo",
    "Estado",
    "Devengado",
    "Deducciones",
    "Neto a pagar",
]


def _filtros_texto(filtros):
    partes = []
    if filtros.get("periodo_id"):
        partes.append(f"Periodo #{filtros['periodo_id']}")
    if filtros.get("estado"):
        partes.append(f"Estado {filtros['estado']}")
    if filtros.get("desde"):
        partes.append(f"desde {_fecha_larga(filtros['desde'])}")
    if filtros.get("hasta"):
        partes.append(f"hasta {_fecha_larga(filtros['hasta'])}")
    return "; ".join(partes) if partes else "Sin filtros (todas las nóminas)"


def _filas(nominas):
    for nomina in nominas:
        empleado = nomina["empleado"]
        periodo = nomina["periodo"]
        yield {
            "id": nomina["id"],
            "empleado": f"{empleado['nombres']} {empleado['apellidos']}",
            "documento": f"{empleado['tipo_documento']} {empleado['numero_documento']}",
            "periodo": f"{periodo['fecha_inicio']} a {periodo['fecha_fin']}",
            "estado": nomina["estado"],
            "total_devengado": nomina["total_devengado"] or 0.0,
            "total_deducciones": nomina["total_deducciones"] or 0.0,
            "neto_pagar": nomina["neto_pagar"] or 0.0,
        }


def _nombre_archivo(empresa, extension):
    base = empresa["razon_social"].replace(" ", "_")
    sello = datetime.now().strftime("%Y%m%d_%H%M")
    return f"reporte_nominas_{base}_{sello}.{extension}"


def build_report_pdf(current_user, periodo_id=None, estado=None, desde=None, hasta=None):
    datos = report_data(current_user, periodo_id, estado, desde, hasta)
    empresa = datos["empresa"]
    resumen = datos["resumen"]
    filas = list(_filas(datos["nominas"]))

    base_styles = getSampleStyleSheet()
    styles = {
        "titulo": ParagraphStyle(
            "titulo",
            parent=base_styles["Title"],
            fontSize=16,
            textColor=colors.HexColor("#1f3b57"),
            spaceAfter=2,
        ),
        "sub": ParagraphStyle("sub", parent=base_styles["Normal"], fontSize=9, leading=12),
        "seccion": ParagraphStyle(
            "seccion",
            parent=base_styles["Heading2"],
            fontSize=11,
            textColor=colors.HexColor("#1f3b57"),
            spaceBefore=10,
            spaceAfter=4,
        ),
        "cell": ParagraphStyle("cell", parent=base_styles["Normal"], fontSize=8.5, leading=11),
        "pie": ParagraphStyle(
            "pie",
            parent=base_styles["Normal"],
            fontSize=8,
            textColor=colors.HexColor("#6b7682"),
        ),
    }

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        title="Reporte general de nóminas",
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
    )

    elementos = [Paragraph("Reporte general de nóminas", styles["titulo"])]
    encabezado = [empresa["razon_social"], f"NIT {empresa['nit']}"]
    elementos.append(Paragraph("<br/>".join(encabezado), styles["sub"]))
    elementos.append(Paragraph(f"<b>Filtros:</b> {_filtros_texto(datos['filtros'])}", styles["sub"]))
    elementos.append(Spacer(1, 6))

    data = [[Paragraph(f"<b>{c}</b>", styles["cell"]) for c in COLUMNAS]]
    for fila in filas:
        data.append(
            [
                Paragraph(str(fila["id"]), styles["cell"]),
                Paragraph(fila["empleado"], styles["cell"]),
                Paragraph(fila["documento"], styles["cell"]),
                Paragraph(fila["periodo"], styles["cell"]),
                Paragraph(fila["estado"], styles["cell"]),
                Paragraph(_money(fila["total_devengado"]), styles["cell"]),
                Paragraph(_money(fila["total_deducciones"]), styles["cell"]),
                Paragraph(_money(fila["neto_pagar"]), styles["cell"]),
            ]
        )
    data.append(
        [
            Paragraph("", styles["cell"]),
            Paragraph("<b>Totales</b>", styles["cell"]),
            Paragraph("", styles["cell"]),
            Paragraph("", styles["cell"]),
            Paragraph("", styles["cell"]),
            Paragraph(f"<b>{_money(resumen['total_devengado'])}</b>", styles["cell"]),
            Paragraph(f"<b>{_money(resumen['total_deducciones'])}</b>", styles["cell"]),
            Paragraph(f"<b>{_money(resumen['neto_pagar'])}</b>", styles["cell"]),
        ]
    )

    anchos = [16 * mm, 55 * mm, 35 * mm, 50 * mm, 24 * mm, 35 * mm, 35 * mm, 35 * mm]
    tabla = Table(data, colWidths=anchos, hAlign="LEFT", repeatRows=1)
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f3b57")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#b8c4d0")),
                ("ALIGN", (5, 0), (-1, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#eef2f6")]),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#dce3ea")),
                ("LINEABOVE", (0, -1), (-1, -1), 0.8, colors.HexColor("#1f3b57")),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    elementos.append(tabla)

    por_estado = resumen["por_estado"]
    if por_estado:
        conteo = ", ".join(f"{estado}: {cantidad}" for estado, cantidad in sorted(por_estado.items()))
        elementos.append(Spacer(1, 6))
        elementos.append(
            Paragraph(f"<b>Nóminas:</b> {resumen['nominas']} ({conteo})", styles["sub"])
        )

    elementos.append(Spacer(1, 12))
    generado = datetime.now().strftime("%d/%m/%Y %H:%M")
    elementos.append(
        Paragraph(
            f"Documento generado el {generado}. Los totales excluyen las nóminas anuladas.",
            styles["pie"],
        )
    )

    doc.build(elementos)
    buffer.seek(0)
    return buffer, _nombre_archivo(empresa, "pdf")


def build_report_excel(current_user, periodo_id=None, estado=None, desde=None, hasta=None):
    datos = report_data(current_user, periodo_id, estado, desde, hasta)
    empresa = datos["empresa"]
    resumen = datos["resumen"]
    filas = list(_filas(datos["nominas"]))

    azul = "1F3B57"
    header_fill = PatternFill("solid", fgColor=azul)
    header_font = Font(color="FFFFFF", bold=True)
    bold = Font(bold=True)
    money_fmt = '"$" #,##0.00'
    thin = Side(style="thin", color="B8C4D0")
    borde = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "Nóminas"

    ws["A1"] = empresa["razon_social"]
    ws["A1"].font = Font(bold=True, size=14, color=azul)
    ws["A2"] = f"NIT {empresa['nit']}"
    ws["A3"] = "Reporte general de nóminas"
    ws["A3"].font = bold
    ws["A4"] = f"Filtros: {_filtros_texto(datos['filtros'])}"
    ws["A5"] = f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}"

    encabezado_fila = 7
    for col, titulo in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=encabezado_fila, column=col, value=titulo)
        celda.fill = header_fill
        celda.font = header_font
        celda.border = borde
        celda.alignment = Alignment(horizontal="center", vertical="center")

    fila_actual = encabezado_fila + 1
    for fila in filas:
        valores = [
            fila["id"],
            fila["empleado"],
            fila["documento"],
            fila["periodo"],
            fila["estado"],
            fila["total_devengado"],
            fila["total_deducciones"],
            fila["neto_pagar"],
        ]
        for col, valor in enumerate(valores, start=1):
            celda = ws.cell(row=fila_actual, column=col, value=valor)
            celda.border = borde
            if col >= 6:
                celda.number_format = money_fmt
        fila_actual += 1

    etiqueta = ws.cell(row=fila_actual, column=2, value="Totales")
    etiqueta.font = bold
    for col, valor in ((6, resumen["total_devengado"]), (7, resumen["total_deducciones"]), (8, resumen["neto_pagar"])):
        celda = ws.cell(row=fila_actual, column=col, value=valor)
        celda.font = bold
        celda.number_format = money_fmt
        celda.border = borde

    anchos = [9, 30, 20, 26, 12, 16, 16, 16]
    for col, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(col)].width = ancho
    ws.freeze_panes = ws.cell(row=encabezado_fila + 1, column=1)

    resumen_ws = wb.create_sheet("Resumen")
    resumen_ws["A1"] = "Resumen por estado"
    resumen_ws["A1"].font = bold
    resumen_ws["A2"] = "Estado"
    resumen_ws["B2"] = "Cantidad"
    resumen_ws["A2"].font = header_font
    resumen_ws["B2"].font = header_font
    resumen_ws["A2"].fill = header_fill
    resumen_ws["B2"].fill = header_fill
    fila_resumen = 3
    for estado_nombre, cantidad in sorted(resumen["por_estado"].items()):
        resumen_ws.cell(row=fila_resumen, column=1, value=estado_nombre)
        resumen_ws.cell(row=fila_resumen, column=2, value=cantidad)
        fila_resumen += 1
    resumen_ws.cell(row=fila_resumen + 1, column=1, value="Total nóminas").font = bold
    resumen_ws.cell(row=fila_resumen + 1, column=2, value=resumen["nominas"]).font = bold
    resumen_ws.column_dimensions["A"].width = 18
    resumen_ws.column_dimensions["B"].width = 12

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer, _nombre_archivo(empresa, "xlsx")
